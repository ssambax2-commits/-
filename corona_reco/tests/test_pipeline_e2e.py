# -*- coding: utf-8 -*-
"""합성 데이터 end-to-end (§9: 15 최초/월별, 16 e2e 스모크).

학습→추천→엑셀2종→DB누적→다음달 피드백 주입→재추천 전 과정을 확인.
"""
import datetime as _dt
import os

import pytest
from openpyxl import load_workbook

from corona_reco import config, pipeline, store as store_mod
from corona_reco.tests import synth

REF = _dt.date(2026, 7, 1)


@pytest.fixture
def fast_ml(monkeypatch):
    """CatBoost 앙상블/iteration 축소로 테스트 가속."""
    monkeypatch.setattr(config, "BAGGING_ENSEMBLE_SIZE", 4)
    fast_params = dict(config.CATBOOST_PARAMS)
    fast_params["iterations"] = 40
    monkeypatch.setattr(config, "CATBOOST_PARAMS", fast_params)
    yield


def _col(ws, name):
    for c in range(1, ws.max_column + 1):
        if ws.cell(row=1, column=c).value == name:
            return c
    raise AssertionError(f"컬럼 없음: {name}")


def _fill_feedback_file(path, month="2026-07"):
    """생성된 피드백 파일의 사용자 시트에 성과를 채워넣는다(숨김시트 보존)."""
    wb = load_workbook(path)
    ws = wb["피드백입력"]
    c_flag = _col(ws, "실제입금여부")
    c_amt = _col(ws, "실제입금액")
    c_date = _col(ws, "입금일자")
    c_act = _col(ws, "활동여부")
    r = 2
    filled = 0
    while ws.cell(row=r, column=2).value not in (None, ""):
        ws.cell(row=r, column=c_act, value="Y")
        if r % 2 == 0:  # 절반은 입금 성공
            ws.cell(row=r, column=c_flag, value="Y")
            ws.cell(row=r, column=c_amt, value=100000)
            ws.cell(row=r, column=c_date, value="2026-07-15")
            filled += 1
        else:
            ws.cell(row=r, column=c_flag, value="N")
        r += 1
    wb.save(path)
    return filled


def test_e2e_first_then_monthly(tmp_path, fast_ml):
    data_dir = str(tmp_path / "data")
    train_csv = str(tmp_path / "train.csv")
    act_csv = str(tmp_path / "act.csv")

    # 합성 데이터: 학습(positive 충분) + 활동
    synth.write_csv(synth.generate(n=2500, positive_rate=0.06, ref_date=REF, seed=1), train_csv)
    synth.write_csv(synth.generate(n=1200, positive_rate=0.05, ref_date=REF, seed=2), act_csv)

    # --- 최초 실행 (피드백 없음) ---
    opt = pipeline.RunOptions(
        mode="first", activity_path=act_csv, training_path=train_csv,
        ref_date=REF, month="2026-07", data_dir=data_dir, include_detail=True)
    res = pipeline.run_pipeline(opt)

    assert os.path.exists(res.report_path)
    assert os.path.exists(res.feedback_path)
    assert len(res.borrowers) > 0
    # ML 활성(positive 충분)
    assert res.model_active is True, res.diagnostics

    # 보고서 시트 구성(v4)
    wb = load_workbook(res.report_path)
    for sheet in ["표지", "1팀 추천리스트", "2팀 추천리스트", "팀별 요약",
                  "제외채권", "법조치추천(참고)", "화해·정상·약속자 관리현황", "진단·검증"]:
        assert sheet in wb.sheetnames, wb.sheetnames

    # 팀 리스트 필수 컬럼(원등급/최종등급/상한 등)
    ws1 = wb["1팀 추천리스트"]
    headers = [c.value for c in ws1[1]]
    for need in ["팀", "부담당자", "고객번호", "대출번호", "원등급", "최종등급",
                 "1개월예상회수액", "3개월예상회수액", "상한하향사유", "추천여부"]:
        assert need in headers, headers

    # 3단계 등급만 존재(S/A/B/C/D 없음), 보류는 추천여부 False
    assert set(res.borrowers["최종등급"].unique()) <= {
        config.TIER_IMMEDIATE, config.TIER_MONTH, config.TIER_HOLD}
    holds = res.borrowers[res.borrowers["최종등급"] == config.TIER_HOLD]
    assert not holds["추천여부"].any()
    # 소액(차주합산 500만↓)은 추천 안 됨
    small = res.borrowers[res.borrowers["소액제외여부"]]
    assert not small["추천여부"].any()

    # DB 누적 확인
    db = store_mod.Store(os.path.join(data_dir, config.SQLITE_FILENAME))
    recos = db.get_recommendations("2026-07")
    assert len(recos) == len(res.borrowers)
    meta = db.get_latest_model_meta()
    assert meta is not None and meta["positive수"] > 0
    db.close()

    # 모델/스냅샷 파일 영속화
    assert os.path.exists(os.path.join(data_dir, config.MODEL_FILENAME))

    # --- 피드백 채우기 ---
    filled = _fill_feedback_file(res.feedback_path, month="2026-07")
    assert filled > 0

    # --- 월별 실행 (피드백 반영) ---
    act2_csv = str(tmp_path / "act2.csv")
    synth.write_csv(synth.generate(n=1200, positive_rate=0.05, ref_date=REF, seed=3), act2_csv)
    opt2 = pipeline.RunOptions(
        mode="monthly", activity_path=act2_csv, prev_feedback_path=res.feedback_path,
        ref_date=REF, month="2026-08", data_dir=data_dir)
    res2 = pipeline.run_pipeline(opt2)

    assert os.path.exists(res2.report_path)
    assert res2.diagnostics.get("총 피드백 관측수", 0) > 0
    assert res2.diagnostics.get("EB 관측수", 0) > 0

    # DB에 2026-08 추천 누적
    db2 = store_mod.Store(os.path.join(data_dir, config.SQLITE_FILENAME))
    recos2 = db2.get_recommendations("2026-08")
    assert len(recos2) > 0
    fb = db2.get_feedback()
    assert len(fb) > 0
    db2.close()


def test_stage2_recovery_and_diag(tmp_path, fast_ml):
    """Stage1×Stage2 예상회수 1M/3M 산출 + fallback 진단 표시."""
    data_dir = str(tmp_path / "data")
    train_csv = str(tmp_path / "train.csv")
    act_csv = str(tmp_path / "act.csv")
    synth.write_csv(synth.generate(n=2500, positive_rate=0.08, ref_date=REF, seed=21), train_csv)
    synth.write_csv(synth.generate(n=1000, positive_rate=0.06, ref_date=REF, seed=22), act_csv)

    opt = pipeline.RunOptions(
        mode="first", activity_path=act_csv, training_path=train_csv,
        ref_date=REF, month="2026-07", data_dir=data_dir, include_detail=True)
    res = pipeline.run_pipeline(opt)

    assert "Stage1 사용" in res.diagnostics
    assert "Stage2 fallback" in res.diagnostics
    assert "calibration 적용" in res.diagnostics
    # 1M/3M 예상회수 분리 산출
    assert "예상회수1M" in res.borrowers.columns
    assert "예상회수3M" in res.borrowers.columns
    reco = res.borrowers[res.borrowers["추천여부"]]
    assert (reco["예상회수1M"] >= 0).all()


def test_randombox_from_result(tmp_path, fast_ml):
    """§3 랜덤박스: 결과 차주풀에서 1,000만↑ 상위20 중 무작위 5, 클릭마다 상이."""
    from corona_reco import grading
    data_dir = str(tmp_path / "data")
    train_csv = str(tmp_path / "train.csv")
    act_csv = str(tmp_path / "act.csv")
    synth.write_csv(synth.generate(n=2000, positive_rate=0.07, ref_date=REF, seed=41), train_csv)
    synth.write_csv(synth.generate(n=1500, positive_rate=0.06, ref_date=REF, seed=42), act_csv)
    opt = pipeline.RunOptions(mode="first", activity_path=act_csv, training_path=train_csv,
                              ref_date=REF, month="2026-07", data_dir=data_dir)
    res = pipeline.run_pipeline(opt)
    p1 = grading.randombox_pick(res.borrowers, seed=1)
    p2 = grading.randombox_pick(res.borrowers, seed=2)
    assert len(p1) <= config.RANDOMBOX_PICK
    if len(p1) == config.RANDOMBOX_PICK:
        assert (p1["동일차주계좌합산원금잔액"] >= config.RANDOMBOX_MIN_MULTI_PRINCIPAL).all()
        assert set(p1["고객번호"]) != set(p2["고객번호"]) or len(p1) < 5


def test_first_run_without_training_falls_back(tmp_path):
    """학습데이터 없이 최초 실행 → 규칙/prior 폴백으로 정상 동작."""
    data_dir = str(tmp_path / "data")
    act_csv = str(tmp_path / "act.csv")
    synth.write_csv(synth.generate(n=300, positive_rate=0.05, ref_date=REF, seed=7), act_csv)

    opt = pipeline.RunOptions(
        mode="first", activity_path=act_csv, training_path=None,
        ref_date=REF, month="2026-07", data_dir=data_dir)
    res = pipeline.run_pipeline(opt)
    assert os.path.exists(res.report_path)
    assert res.model_active is False  # ML 폴백
    assert len(res.borrowers) > 0


def test_low_positive_ml_fallback(tmp_path, fast_ml):
    """positive 표본 부족 → ML 비활성(규칙/prior 폴백)."""
    data_dir = str(tmp_path / "data")
    train_csv = str(tmp_path / "train.csv")
    act_csv = str(tmp_path / "act.csv")
    # positive_rate 매우 낮게 → positive < 임계
    synth.write_csv(synth.generate(n=500, positive_rate=0.01, ref_date=REF, seed=11), train_csv)
    synth.write_csv(synth.generate(n=300, positive_rate=0.02, ref_date=REF, seed=12), act_csv)

    opt = pipeline.RunOptions(
        mode="first", activity_path=act_csv, training_path=train_csv,
        ref_date=REF, month="2026-07", data_dir=data_dir)
    res = pipeline.run_pipeline(opt)
    assert res.model_active is False
    assert any("폴백" in w or "임계" in w for w in res.warnings)
