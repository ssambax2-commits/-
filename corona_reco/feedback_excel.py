# -*- coding: utf-8 -*-
"""
feedback_excel.py — 피드백용 엑셀 (다음 달 재업로드용, §5-8·§13)
=================================================================
- 사용자 입력 시트: 전체 채점 차주 수록(추천/비추천 모두 — 자연입금 관측으로
  선택편향 완화). 추천여부/원등급/최종등급/추천순위 저장, 사용자는
  활동여부/실제입금여부/실제입금액/입금일자/비고만 채운다.
- 숨김 시트(모델용): 추천시점 피처벡터, 대출번호(계좌키), 모델버전.
"""
from __future__ import annotations

from typing import Dict

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="44546A")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")

USER_COLUMNS = [
    "추천월", "고객번호", "이름", "원금잔액", "다중계좌원금잔액", "부담당자", "팀",
    "원등급", "최종등급", "추천여부", "정원초과편입여부", "추천순위", "회생",
    "추천사유", "추천점수",
    "활동여부", "실제입금여부", "실제입금액", "입금일자", "비고",
]
INPUT_COLUMNS = {"활동여부", "실제입금여부", "실제입금액", "입금일자", "비고"}
AMOUNT_COLUMNS = {"원금잔액", "다중계좌원금잔액", "실제입금액"}
SCORE_COLUMNS = {"추천점수"}
TEXT_COLUMNS = {"고객번호"}

HIDDEN_COLUMNS = ["추천월", "고객번호", "성명", "대출번호", "최종등급",
                  "모델버전", "피처벡터"]


def _autosize(ws, headers, rows):
    for c, h in enumerate(headers, start=1):
        w = sum(2 if ord(ch) > 0x1100 else 1 for ch in str(h))
        for r in rows[:200]:
            v = r[c - 1]
            if v is not None and str(v) != "":
                w = max(w, sum(2 if ord(ch) > 0x1100 else 1 for ch in str(v)[:40]))
        ws.column_dimensions[get_column_letter(c)].width = min(max(w + 2, 8), 46)


def write_feedback_file(path: str, borrowers: pd.DataFrame,
                        feature_json_by_key: Dict[str, str],
                        model_version: str, month: str) -> str:
    """전체 채점 차주를 담은 피드백 파일 저장."""
    wb = Workbook()
    ws = wb.active
    ws.title = "피드백입력"

    for c, h in enumerate(USER_COLUMNS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")

    rows = []
    for _, r in borrowers.iterrows():
        rows.append([
            month,
            str(r.get("고객번호", "")),
            r.get("성명", ""),
            float(r.get("원금잔액", 0) or 0),
            float(r.get("다중계좌원금잔액", 0) or 0),
            r.get("부담당자", ""),
            r.get("팀", ""),
            str(r.get("원등급", "")),
            str(r.get("최종등급", "")),
            "Y" if bool(r.get("추천여부", False)) else "N",
            "Y" if bool(r.get("정원초과편입여부", False)) else "",
            r.get("추천순위_차주기준"),
            r.get("회생", ""),
            r.get("추천사유", ""),
            round(float(r.get("borrower_score", 0) or 0), 2),
            "", "", "", "", "",   # 사용자 입력칸
        ])

    for ri, row in enumerate(rows, start=2):
        for c, h in enumerate(USER_COLUMNS, start=1):
            cell = ws.cell(row=ri, column=c, value=row[c - 1])
            if h in AMOUNT_COLUMNS:
                cell.number_format = "#,##0"
            if h in SCORE_COLUMNS:
                cell.number_format = "0.00"
            if h in TEXT_COLUMNS:
                cell.number_format = "@"
            if h in INPUT_COLUMNS:
                cell.fill = _INPUT_FILL

    ws.freeze_panes = "A2"
    if rows:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(USER_COLUMNS))}{len(rows) + 1}"
    _autosize(ws, USER_COLUMNS, rows)
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"

    # 숨김 시트(모델용)
    wsh = wb.create_sheet("_모델메타")
    for c, h in enumerate(HIDDEN_COLUMNS, start=1):
        wsh.cell(row=1, column=c, value=h)
    for ri, (_, r) in enumerate(borrowers.iterrows(), start=2):
        loan_no = str(r.get("대표대출번호", ""))
        vals = [month, str(r.get("고객번호", "")), r.get("성명", ""), loan_no,
                str(r.get("최종등급", "")), model_version,
                feature_json_by_key.get(loan_no, "")]
        for c, v in enumerate(vals, start=1):
            wsh.cell(row=ri, column=c, value=v)
    wsh.sheet_state = "hidden"

    wb.save(path)
    return path
